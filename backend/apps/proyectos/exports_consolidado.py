"""
Exportacion del informe consolidado a PDF y Excel (HU-06, T-107 y T-108).

Recibe el diccionario que arma services_consolidado.consolidar() y lo
convierte en un archivo descargable. Aqui no se consulta la base de datos:
todo el calculo ya viene hecho, de modo que agregar una columna al informe no
obliga a tocar dos veces la misma logica.

Ambas funciones devuelven un BytesIO posicionado al inicio, listo para
adjuntar a un FileResponse.

Funciones:
- exportar_consolidado_pdf(informe, titulo): documento reportlab apaisado con
  el resumen, el presupuesto, las distribuciones y la tabla de proyectos.
- exportar_consolidado_excel(informe, titulo): libro openpyxl con una hoja
  por bloque, para que el Coordinador RSU pueda seguir analizando los datos.

Sigue el mismo patron que apps/planificacion/services.py, que exporta la
matriz operativa.

Conecta con:
- apps/proyectos/services_consolidado.py: origen de los datos.
- apps/proyectos/views_consolidado.py: vistas que descargan estos archivos.
"""
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
)

# Paleta institucional, compartida por el PDF y el Excel para que ambos
# documentos se vean como parte del mismo informe.
AZUL = '1F4E79'
AZUL_CLARO = 'DCE6F1'
GRIS = 'F2F2F2'

ETIQUETAS_FILTRO = {
    'facultad': 'Facultad',
    'escuela': 'Escuela profesional',
    'departamento': 'Departamento academico',
    'eje_rsu': 'Eje RSU',
    'ods': 'ODS',
    'periodo': 'Periodo academico',
    'estado': 'Estado',
}


def _fecha_legible(iso):
    """Convierte una marca de tiempo ISO a un texto corto para la cabecera."""
    if not iso:
        return ''
    try:
        return datetime.fromisoformat(iso).strftime('%d/%m/%Y %H:%M')
    except ValueError:
        return iso


def _texto_filtros(informe):
    """Resume en una linea los filtros con los que se genero el informe."""
    filtros = informe.get('filtros_aplicados') or {}
    if not filtros:
        return 'Sin filtros: todos los proyectos aprobados y finalizados visibles.'
    partes = [
        '%s = %s' % (ETIQUETAS_FILTRO.get(k, k), v)
        for k, v in filtros.items()
    ]
    return 'Filtros aplicados: ' + '; '.join(partes)


# ─────────────────────────────────────────────────────────────────────────────
# T-107: PDF
# ─────────────────────────────────────────────────────────────────────────────

def _estilos_pdf():
    base = getSampleStyleSheet()
    return {
        'titulo': ParagraphStyle(
            'TituloRSU', parent=base['Title'], fontSize=16,
            textColor=colors.HexColor('#' + AZUL), spaceAfter=6),
        'subtitulo': ParagraphStyle(
            'SubtituloRSU', parent=base['Normal'], fontSize=9,
            textColor=colors.HexColor('#555555'), spaceAfter=2),
        'seccion': ParagraphStyle(
            'SeccionRSU', parent=base['Heading2'], fontSize=12,
            textColor=colors.HexColor('#' + AZUL), spaceBefore=12, spaceAfter=6),
        'celda': ParagraphStyle(
            'CeldaRSU', parent=base['Normal'], fontSize=7, leading=9),
        'normal': base['Normal'],
    }


def _estilo_tabla(con_cabecera=True):
    reglas = [
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#BFBFBF')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('FONTSIZE', (0, 0), (-1, -1), 7),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#' + GRIS)]),
    ]
    if con_cabecera:
        reglas += [
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#' + AZUL)),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ]
    return TableStyle(reglas)


def _tabla_indicadores(informe):
    """Tabla de dos columnas con los numeros gruesos del informe."""
    resumen = informe['resumen']
    presupuesto = informe['presupuesto']
    metas = informe['metas']
    filas = [
        ['Indicador', 'Valor'],
        ['Proyectos consolidados', resumen['total_proyectos']],
        ['   Aprobados', resumen['aprobados']],
        ['   Finalizados', resumen['finalizados']],
        ['Docentes responsables', resumen['docentes_responsables']],
        ['Docentes declarados', resumen['total_docentes_declarados']],
        ['Estudiantes declarados', resumen['total_estudiantes_declarados']],
        ['Avance promedio de ejecucion', '%s %%' % resumen['avance_promedio']],
        ['Actividades completadas',
         '%s de %s (%s %%)' % (resumen['actividades_completadas'],
                               resumen['actividades_total'],
                               resumen['porcentaje_actividades_completadas'])],
        ['Presupuesto financiado', 'S/ %.2f' % presupuesto['monto_financiado']],
        ['Presupuesto programado', 'S/ %.2f' % presupuesto['monto_presupuestado']],
        ['Presupuesto ejecutado', 'S/ %.2f' % presupuesto['monto_ejecutado']],
        ['Saldo por ejecutar', 'S/ %.2f' % presupuesto['saldo_por_ejecutar']],
        ['Ejecucion presupuestal',
         '%s %%' % presupuesto['porcentaje_ejecucion_presupuestal']],
        ['Metas cumplidas',
         '%s de %s (%s %%)' % (metas['cumplidas'], metas['total'],
                               metas['porcentaje_cumplimiento'])],
    ]
    tabla = Table(filas, colWidths=[110 * mm, 60 * mm])
    tabla.setStyle(_estilo_tabla())
    return tabla


def _tabla_distribucion(titulo, filas, estilos):
    """Tabla generica de distribucion: etiqueta y cantidad de proyectos."""
    datos = [[titulo, 'Proyectos']]
    for fila in filas:
        datos.append([Paragraph(str(fila['etiqueta']), estilos['celda']), fila['total']])
    if len(datos) == 1:
        datos.append(['Sin datos', 0])
    tabla = Table(datos, colWidths=[130 * mm, 40 * mm])
    tabla.setStyle(_estilo_tabla())
    return tabla


def _tabla_proyectos(informe, estilos):
    """Tabla principal: una fila por proyecto consolidado."""
    cabecera = ['Codigo', 'Titulo', 'Facultad', 'Eje RSU', 'Periodo', 'Estado',
                'Avance', 'Programado S/', 'Ejecutado S/', 'Metas']
    datos = [cabecera]
    for p in informe.get('proyectos', []):
        datos.append([
            Paragraph(p['codigo'], estilos['celda']),
            Paragraph(p['titulo'], estilos['celda']),
            Paragraph(p['facultad'] or '', estilos['celda']),
            Paragraph(p['eje_rsu'] or '', estilos['celda']),
            Paragraph(p['periodo'] or '', estilos['celda']),
            Paragraph(p['estado_display'], estilos['celda']),
            '%s %%' % p['avance']['porcentaje_ejecucion'],
            '%.2f' % p['presupuesto']['monto_presupuestado'],
            '%.2f' % p['presupuesto']['monto_ejecutado'],
            '%s/%s' % (p['metas']['cumplidas'], p['metas']['total']),
        ])
    if len(datos) == 1:
        datos.append(['', Paragraph('No hay proyectos que cumplan los filtros.',
                                    estilos['celda']), '', '', '', '', '', '', '', ''])

    anchos = [22 * mm, 62 * mm, 34 * mm, 30 * mm, 20 * mm, 20 * mm,
              16 * mm, 22 * mm, 22 * mm, 15 * mm]
    tabla = Table(datos, colWidths=anchos, repeatRows=1)
    tabla.setStyle(_estilo_tabla())
    return tabla


def exportar_consolidado_pdf(informe, titulo='Informe consolidado de proyectos RSU'):
    """Genera el PDF del informe consolidado y devuelve un BytesIO."""
    salida = io.BytesIO()
    doc = SimpleDocTemplate(
        salida, pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm,
        title=titulo, author='Sistema Web de Gestion RSU - UNSA',
    )
    estilos = _estilos_pdf()
    elementos = [
        Paragraph(titulo, estilos['titulo']),
        Paragraph('Universidad Nacional de San Agustin de Arequipa - Oficina de '
                  'Responsabilidad Social Universitaria', estilos['subtitulo']),
        Paragraph('Generado el %s' % _fecha_legible(informe.get('generado_en')),
                  estilos['subtitulo']),
        Paragraph('Alcance: proyectos en estado aprobado y finalizado.',
                  estilos['subtitulo']),
        Paragraph(_texto_filtros(informe), estilos['subtitulo']),
        Spacer(1, 8),

        Paragraph('Indicadores consolidados', estilos['seccion']),
        _tabla_indicadores(informe),

        Paragraph('Distribucion por facultad', estilos['seccion']),
        _tabla_distribucion('Facultad', informe['distribuciones']['por_facultad'], estilos),

        Paragraph('Distribucion por eje RSU', estilos['seccion']),
        _tabla_distribucion('Eje RSU', informe['distribuciones']['por_eje_rsu'], estilos),

        Paragraph('Distribucion por ODS', estilos['seccion']),
        _tabla_distribucion('Objetivo de Desarrollo Sostenible',
                            informe['distribuciones']['por_ods'], estilos),

        Paragraph('Distribucion por periodo academico', estilos['seccion']),
        _tabla_distribucion('Periodo', informe['distribuciones']['por_periodo'], estilos),

        PageBreak(),
        Paragraph('Detalle de proyectos', estilos['seccion']),
        _tabla_proyectos(informe, estilos),
    ]
    doc.build(elementos)
    salida.seek(0)
    return salida


# ─────────────────────────────────────────────────────────────────────────────
# T-108: Excel
# ─────────────────────────────────────────────────────────────────────────────

_LINEA = Side(style='thin', color='BFBFBF')
_BORDE = Border(left=_LINEA, right=_LINEA, top=_LINEA, bottom=_LINEA)
_FUENTE_CABECERA = Font(name='Calibri', size=10, bold=True, color='FFFFFF')
_FUENTE_TITULO = Font(name='Calibri', size=13, bold=True, color=AZUL)
_FUENTE_SECCION = Font(name='Calibri', size=11, bold=True, color=AZUL)
_FUENTE_NORMAL = Font(name='Calibri', size=10)
_RELLENO_CABECERA = PatternFill('solid', fgColor=AZUL)
_RELLENO_SECCION = PatternFill('solid', fgColor=AZUL_CLARO)
_AJUSTE = Alignment(vertical='top', wrap_text=True)


def _escribir_cabecera(ws, fila, columnas, anchos=None):
    """Escribe la fila de encabezados y deja la hoja lista para los datos."""
    for i, nombre in enumerate(columnas, 1):
        celda = ws.cell(row=fila, column=i, value=nombre)
        celda.font = _FUENTE_CABECERA
        celda.fill = _RELLENO_CABECERA
        celda.alignment = _AJUSTE
        celda.border = _BORDE
    if anchos:
        for i, ancho in enumerate(anchos, 1):
            ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)
    return fila + 1


def _escribir_fila(ws, fila, valores):
    for i, valor in enumerate(valores, 1):
        celda = ws.cell(row=fila, column=i, value=valor)
        celda.font = _FUENTE_NORMAL
        celda.alignment = _AJUSTE
        celda.border = _BORDE
    return fila + 1


def _hoja_resumen(wb, informe, titulo):
    ws = wb.active
    ws.title = 'Resumen'
    ws.column_dimensions['A'].width = 44
    ws.column_dimensions['B'].width = 30

    ws['A1'] = titulo
    ws['A1'].font = _FUENTE_TITULO
    ws['A2'] = 'Universidad Nacional de San Agustin de Arequipa - Oficina de RSU'
    ws['A3'] = 'Generado el %s' % _fecha_legible(informe.get('generado_en'))
    ws['A4'] = 'Alcance: proyectos en estado aprobado y finalizado.'
    ws['A5'] = _texto_filtros(informe)
    for ref in ('A2', 'A3', 'A4', 'A5'):
        ws[ref].font = Font(name='Calibri', size=9, italic=True, color='555555')

    resumen = informe['resumen']
    presupuesto = informe['presupuesto']
    metas = informe['metas']

    bloques = [
        ('Proyectos', [
            ('Proyectos consolidados', resumen['total_proyectos']),
            ('Aprobados', resumen['aprobados']),
            ('Finalizados', resumen['finalizados']),
            ('Avance promedio de ejecucion (%)', resumen['avance_promedio']),
        ]),
        ('Participacion', [
            ('Docentes responsables distintos', resumen['docentes_responsables']),
            ('Docentes declarados', resumen['total_docentes_declarados']),
            ('Estudiantes declarados', resumen['total_estudiantes_declarados']),
        ]),
        ('Actividades', [
            ('Actividades registradas', resumen['actividades_total']),
            ('Actividades completadas', resumen['actividades_completadas']),
            ('Completadas (%)', resumen['porcentaje_actividades_completadas']),
        ]),
        ('Presupuesto (S/)', [
            ('Financiado', presupuesto['monto_financiado']),
            ('Programado', presupuesto['monto_presupuestado']),
            ('Ejecutado', presupuesto['monto_ejecutado']),
            ('Saldo por ejecutar', presupuesto['saldo_por_ejecutar']),
            ('Ejecucion presupuestal (%)',
             presupuesto['porcentaje_ejecucion_presupuestal']),
        ]),
        ('Metas e indicadores', [
            ('Metas registradas', metas['total']),
            ('Metas cumplidas', metas['cumplidas']),
            ('Cumplimiento (%)', metas['porcentaje_cumplimiento']),
        ]),
    ]

    fila = 7
    for nombre, filas in bloques:
        celda = ws.cell(row=fila, column=1, value=nombre)
        celda.font = _FUENTE_SECCION
        celda.fill = _RELLENO_SECCION
        ws.cell(row=fila, column=2).fill = _RELLENO_SECCION
        fila += 1
        for etiqueta, valor in filas:
            fila = _escribir_fila(ws, fila, [etiqueta, valor])
        fila += 1


def _hoja_proyectos(wb, informe):
    ws = wb.create_sheet('Proyectos')
    columnas = ['Codigo', 'Titulo', 'Estado', 'Facultad', 'Escuela', 'Departamento',
                'Periodo', 'Semestre', 'Eje RSU', 'ODS', 'Docente responsable',
                'Docentes', 'Estudiantes', 'Avance (%)', 'Financiado S/',
                'Programado S/', 'Ejecutado S/', 'Ejecucion presupuestal (%)',
                'Metas', 'Metas cumplidas', 'Cumplimiento metas (%)',
                'Actividades', 'Actividades completadas', 'Avances', 'Evidencias',
                'Fecha inicio', 'Fecha termino', 'Fecha aprobacion']
    anchos = [16, 52, 14, 30, 30, 30, 14, 12, 26, 26, 30, 10, 12, 11, 14, 14, 14,
              16, 8, 14, 16, 12, 16, 10, 11, 13, 13, 15]
    fila = _escribir_cabecera(ws, 1, columnas, anchos)

    for p in informe.get('proyectos', []):
        fila = _escribir_fila(ws, fila, [
            p['codigo'], p['titulo'], p['estado_display'],
            p['facultad'], p['escuela'], p['departamento'],
            p['periodo'], p['semestre_academico'], p['eje_rsu'],
            ', '.join('ODS %s' % o['numero'] for o in p['ods']),
            p['docente_responsable'], p['nro_docentes'], p['nro_estudiantes'],
            p['avance']['porcentaje_ejecucion'],
            p['presupuesto']['monto_financiado'],
            p['presupuesto']['monto_presupuestado'],
            p['presupuesto']['monto_ejecutado'],
            p['presupuesto']['porcentaje_ejecucion_presupuestal'],
            p['metas']['total'], p['metas']['cumplidas'],
            p['metas']['porcentaje_cumplimiento'],
            p['avance']['actividades_total'], p['avance']['actividades_completadas'],
            p['avance']['avances_registrados'], p['avance']['evidencias_vigentes'],
            p['fecha_inicio'], p['fecha_termino'], p['fecha_aprobacion'],
        ])
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(columnas)), max(fila - 1, 1))


def _hoja_distribuciones(wb, informe):
    ws = wb.create_sheet('Distribuciones')
    ws.column_dimensions['A'].width = 22
    ws.column_dimensions['B'].width = 58
    ws.column_dimensions['C'].width = 14

    fila = _escribir_cabecera(ws, 1, ['Dimension', 'Etiqueta', 'Proyectos'],
                              [22, 58, 14])
    dimensiones = [
        ('Facultad', informe['distribuciones']['por_facultad']),
        ('Escuela profesional', informe['distribuciones']['por_escuela']),
        ('Eje RSU', informe['distribuciones']['por_eje_rsu']),
        ('ODS', informe['distribuciones']['por_ods']),
        ('Periodo academico', informe['distribuciones']['por_periodo']),
        ('Estado', informe['distribuciones']['por_estado']),
    ]
    for nombre, filas in dimensiones:
        for f in filas:
            fila = _escribir_fila(ws, fila, [nombre, f['etiqueta'], f['total']])


def exportar_consolidado_excel(informe, titulo='Informe consolidado de proyectos RSU'):
    """Genera el libro Excel del informe consolidado y devuelve un BytesIO."""
    wb = openpyxl.Workbook()
    _hoja_resumen(wb, informe, titulo)
    _hoja_proyectos(wb, informe)
    _hoja_distribuciones(wb, informe)

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida
